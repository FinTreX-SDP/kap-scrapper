"""Where the data goes: the database path, an Excel summary and per disclosure a Markdown and a JSON file.

The database is the full record. The Excel file and the Markdown detail files are for people to look at:
the Excel file lists the newest disclosures with their main fields, and each row links to the
disclosure's detail file, which shows everything extracted from it. The JSON file holds the same
content as data, for programs such as a model that reads disclosures.

Usage:
    python kap_output.py   # rebuild the Excel summary and every detail and JSON file from the database
"""
import json
import re
from itertools import groupby
from operator import itemgetter
from pathlib import Path

import duckdb
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

DATA_DIR = Path(__file__).parent / "data"
DB_PATH = DATA_DIR / "kap.duckdb"
EXCEL_PATH = DATA_DIR / "disclosures.xlsx"
DETAILS_DIR = DATA_DIR / "details"
JSON_DIR = DATA_DIR / "json"
EXCEL_MAX_ROWS = 5000  # newest disclosures in the Excel summary; the database keeps all of them

EXCEL_COLUMNS = {  # header -> column width
    "disclosure_index": 12, "publish_date": 19, "stock_code": 12, "company": 40, "title": 45,
    "summary": 60, "class": 7, "attachments": 12, "attachments_read": 16, "financial_values": 16,
    "kap": 6, "detail_file": 28,
}


def table_exists(con: duckdb.DuckDBPyConnection, name: str) -> bool:
    return con.execute("SELECT count(*) FROM duckdb_tables() WHERE table_name = ?", [name]).fetchone()[0] > 0


def optional_rows(con: duckdb.DuckDBPyConnection, table: str, query: str, params: list) -> list[dict]:
    """Rows of a query on a table that may not have been created yet, keyed by column name."""
    if not table_exists(con, table):
        return []
    cursor = con.execute(query, params)
    names = [column[0] for column in cursor.description]
    return [dict(zip(names, row)) for row in cursor.fetchall()]


def detail_path(disclosure_index: int, published, stock_code: str | None) -> Path:
    stock = re.sub(r"[^A-Za-z0-9]+", "-", stock_code or "").strip("-")[:40]
    name = f"{disclosure_index}_{stock}.md" if stock else f"{disclosure_index}.md"
    return DETAILS_DIR / f"{published:%Y-%m-%d}" / name


def excel_value(value):
    return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value


def export_to_excel(con: duckdb.DuckDBPyConnection) -> int:
    """Write the newest disclosures to the Excel summary; return the number of disclosures in the database."""
    read = ("(SELECT count(*) FROM attachment_texts t WHERE t.disclosure_index = d.disclosure_index)"
            if table_exists(con, "attachment_texts") else "NULL")
    values = ("(SELECT NULLIF(count(*), 0) FROM financial_items f WHERE f.disclosure_index = d.disclosure_index)"
              if table_exists(con, "financial_items") else "NULL")
    rows = con.execute(f"""
        SELECT d.disclosure_index, d.publish_date, d.stock_code, d.company_title, d.title, d.summary,
               d.disclosure_class, d.attachment_count, {read}, {values}, d.url
        FROM disclosures d ORDER BY d.publish_date DESC LIMIT {EXCEL_MAX_ROWS}
    """).fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = "Disclosures"
    ws.append(list(EXCEL_COLUMNS))
    for *fields, url in rows:
        detail = detail_path(fields[0], fields[1], fields[2])
        ws.append([excel_value(v) for v in fields] + ["KAP", detail.name if detail.exists() else None])
        link = ws.cell(ws.max_row, len(EXCEL_COLUMNS) - 1)
        link.hyperlink, link.style = url, "Hyperlink"
        if detail.exists():  # relative to the workbook, so the data folder can be moved
            link = ws.cell(ws.max_row, len(EXCEL_COLUMNS))
            link.hyperlink, link.style = detail.relative_to(DATA_DIR).as_posix(), "Hyperlink"
    for letter, width in zip("ABCDEFGHIJKL", EXCEL_COLUMNS.values()):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(EXCEL_PATH)
    return con.execute("SELECT count(*) FROM disclosures").fetchone()[0]


def md_cell(value) -> str:
    return "" if value is None else str(value).replace("|", "\\|").replace("\n", "<br>")


def md_table(header: list, rows: list[list]) -> list[str]:
    width = max([len(header)] + [len(r) for r in rows])
    lines = ["| " + " | ".join(md_cell(v) for v in list(header) + [""] * (width - len(header))) + " |",
             "|" + "---|" * width]
    lines += ["| " + " | ".join(md_cell(v) for v in list(r) + [""] * (width - len(r))) + " |" for r in rows]
    return lines + [""]


def fmt_number(value: float | None) -> str:
    """Turkish number format, as on KAP: 1.234.567,89"""
    if value is None:
        return ""
    text = f"{value:,.2f}" if value % 1 else f"{value:,.0f}"
    return text.replace(",", "_").replace(".", ",").replace("_", ".")


def statement_table(items: list[dict]) -> list[str]:
    """One financial statement as a table: items as rows and periods (or equity components) as columns."""
    columns, rows = [], {}
    for item in items:
        member, period = item["member"], item["period_label"]
        column = member or period  # changes in equity: components are the columns, periods group the rows
        if column not in columns:
            columns.append(column)
        rows.setdefault((period if member else None, item["concept"], item["label"]), {})[column] = item["value"]
    grouped = any(key[0] for key in rows)
    header = (["Period"] if grouped else []) + ["Item"] + columns
    body = [([key[0]] if grouped else []) + [key[2]] + [fmt_number(cells.get(c)) for c in columns]
            for key, cells in rows.items()]
    return md_table(header, body)


STATEMENT_KEYS = ("statement", "currency", "consolidation")


def disclosure_record(con: duckdb.DuckDBPyConnection, disclosure_index: int) -> dict:
    """Everything extracted from one disclosure so far, as plain data. Dates stay date objects."""
    params = [disclosure_index]
    record = optional_rows(con, "disclosures", "SELECT * FROM disclosures WHERE disclosure_index = ?", params)[0]
    page = optional_rows(con, "disclosure_details", """
        SELECT page_format, text FROM disclosure_details WHERE disclosure_index = ?
    """, params)
    page = page[0] if page else {"page_format": None, "text": None}

    financials = optional_rows(con, "financial_items", """
        SELECT statement, currency, consolidation, concept, label, member, period_label, period_start, period_end, value
        FROM financial_items WHERE disclosure_index = ? ORDER BY statement, row_no
    """, params)
    statements = [{**dict(zip(STATEMENT_KEYS, key)),
                   "items": [{k: v for k, v in item.items() if k not in STATEMENT_KEYS} for item in group]}
                  for key, group in groupby(financials, key=itemgetter(*STATEMENT_KEYS))]
    # A financial report's fields are its statements, so they are left out to avoid storing them twice.
    fields = [] if financials else optional_rows(con, "disclosure_fields", """
        SELECT section, concept, label, value FROM disclosure_fields WHERE disclosure_index = ? ORDER BY row_no, value_col
    """, params)
    tables = [{"section": t["section"], "rows": json.loads(t["rows_json"])} for t in optional_rows(con, "disclosure_tables", """
        SELECT section, rows_json FROM disclosure_tables WHERE disclosure_index = ? ORDER BY table_no
    """, params)]

    texts = {t["file_id"]: t for t in optional_rows(con, "attachment_texts", """
        SELECT file_id, file_type, page_count, ocr_pages, skipped_pages, text FROM attachment_texts
        WHERE disclosure_index = ?
    """, params)}
    attachments = optional_rows(con, "disclosure_attachments", """
        SELECT file_id, file_name, url FROM disclosure_attachments WHERE disclosure_index = ? ORDER BY file_name
    """, params)
    for attachment in attachments:
        if attachment["file_id"] not in texts:
            continue  # not processed yet: no text or tables
        attachment |= texts[attachment["file_id"]]
        table_rows = optional_rows(con, "attachment_tables", """
            SELECT table_no, page, method, cells FROM attachment_tables
            WHERE disclosure_index = ? AND file_id = ? ORDER BY table_no, row_no
        """, [disclosure_index, attachment["file_id"]])
        attachment["tables"] = [{"page": rows[0]["page"], "method": rows[0]["method"], "rows": [r["cells"] for r in rows]}
                                for rows in (list(g) for _, g in groupby(table_rows, key=itemgetter("table_no")))]

    return record | {"page_format": page["page_format"], "fields": fields, "tables": tables,
                     "financial_statements": statements, "attachments": attachments, "text": page["text"]}


def write_detail_file(con: duckdb.DuckDBPyConnection, disclosure_index: int) -> Path:
    """Write one disclosure, with everything extracted from it so far, to a Markdown file and a JSON file."""
    r = disclosure_record(con, disclosure_index)
    lines = [f"# {r['stock_code'] or '-'}: {r['title']}", "",
             f"- **Company:** {r['company_title'] or '-'}", f"- **Published:** {r['publish_date']}",
             f"- **Class:** {r['disclosure_class']}", f"- **Summary:** {r['summary'] or '-'}", f"- **KAP:** {r['url']}", ""]

    if r["fields"]:
        lines += ["## Fields", ""]
        for section, group in groupby(r["fields"], key=itemgetter("section")):
            lines += ([f"### {section}", ""] if section else []) + md_table(["Label", "Value"],
                                                                            [[f["label"], f["value"]] for f in group])

    if r["tables"]:
        lines += ["## Tables on the page", ""]
        for n, table in enumerate(r["tables"], 1):
            rows = table["rows"]
            lines += [f"### {table['section'] or f'Table {n}'}", ""] + md_table(rows[0] if rows else [], rows[1:])

    if r["financial_statements"]:
        lines += ["## Financial statements", ""]
        for statement in r["financial_statements"]:
            lines += [f"### {statement['statement']}", ""] + statement_table(statement["items"])

    if r["attachments"]:
        lines += ["## Attachments", ""]
    for a in r["attachments"]:
        lines += [f"### {a['file_name']}", "", f"Download: {a['url']}", ""]
        if "text" not in a:
            lines += ["Not processed yet.", ""]
            continue
        lines += [f"{a['file_type']}, {a['page_count']} pages, {a['ocr_pages']} read with OCR, "
                  f"{a['skipped_pages']} skipped.", ""]
        for n, table in enumerate(a["tables"], 1):
            rows = table["rows"]
            lines += [f"**Table {n}** (page {table['page']}, {table['method']})", ""] + md_table(rows[0], rows[1:])
        if a["text"]:
            lines += ["````text", a["text"], "````", ""]

    if r["text"]:
        lines += ["## Page text", "", "````text", r["text"], "````", ""]

    path = detail_path(disclosure_index, r["publish_date"], r["stock_code"])
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")
    json_path = JSON_DIR / path.relative_to(DETAILS_DIR).with_suffix(".json")
    json_path.parent.mkdir(parents=True, exist_ok=True)
    json_path.write_text(json.dumps(r, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    return path


def main() -> None:
    con = duckdb.connect(str(DB_PATH))
    indexes = [r["disclosure_index"] for r in optional_rows(con, "disclosure_details",
                                                            "SELECT disclosure_index FROM disclosure_details", [])]
    for disclosure_index in indexes:
        write_detail_file(con, disclosure_index)
    total = export_to_excel(con)
    con.close()
    print(f"Detail and JSON files written: {len(indexes)}")
    print(f"Excel summary: newest {min(total, EXCEL_MAX_ROWS)} of {total} disclosures, {EXCEL_PATH}")


if __name__ == "__main__":
    main()
